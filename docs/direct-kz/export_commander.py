#!/usr/bin/env python3
"""Файлы для импорта в Директ Коммандер: кампании Praxen AI на Казахстан (план — docs/DIRECT-KZ.md).

Формат и приёмы — как в проекте WinWinBot (docs/marketing/direct/export_commander.py), где файлы уже
проходили импорт: официальный шаблон Яндекса `template/direct_example.xlsx`, справка Коммандера
«Формат файла». XLSX = одна кампания; лист «Тексты»: E7 тип кампании, E9 минус-фразы на кампанию,
данные с 12-й строки. Тип кампании «Единая перфоманс-кампания», объявление «Комбинаторное»:
до 7 заголовков (≤56, слово ≤22) и 3 текстов (≤81, ≤15 знаков препинания, слово ≤23).
Перед блоком «Комбинаторика» — 3 пустые старые колонки «Заголовок 1; Заголовок 2; Текст»: без них
Коммандер относит комбинаторные «Заголовок 1/2» к старым полям и теряет их.

Бюджет, стратегия, места показа, ставки и картинки в файл импорта не входят — они в
`commander/Настройки_кампаний_KZ.xlsx` и задаются в Коммандере после импорта.
Запуск: python3 -I docs/direct-kz/export_commander.py
"""
import re
import sys
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "template" / "direct_example.xlsx"
OUT = HERE / "commander"
SITE = "https://praxenai.ge"
FIRST_ROW = 12
LEGACY = ["Заголовок 1", "Заголовок 2", "Текст"]
SHIFT = len(LEGACY)
C = {"plus": 1, "type": 2, "group": 4, "num": 5, "phrase": 7, "h": 9, "t": 16, "len": 19, "url": 41,
     "display": 42, "region": 43, "sl_t": 49, "sl_d": 50, "sl_u": 51, "callouts": 57}
C = {k: v + SHIFT if v >= 9 else v for k, v in C.items()}
LEN_COLS = [get_column_letter(C["h"] + i) for i in range(10)]
PUNCT = re.compile(r"[.,!?;:\-—–()«»\"']")
REGION = "Казахстан"


def utm(slug, path, term=True):
    tail = "&utm_content={ad_id}&utm_term={keyword}" if term else ""
    q = f"utm_source=yandex&utm_medium=cpc&utm_campaign={slug}{tail}"
    base, _, frag = path.partition("#")
    return f"{SITE}{base}?{q}" + (f"#{frag}" if frag else "")


# Минус-слова: общие + свои для кампании. «чат» не минусуем — им заблокировались бы фразы про чат-ботов.
MINUS_COMMON = ("рулетка рулетки знакомства реферат курсовая диплом егэ огэ гдз школьник детей дети игра игры "
                "вакансия вакансии работа удаленно зарплата скачать apk взлом nsfw 18 картинки картинка нарисовать "
                "фото видео музыка песня голос озвучка конструктор самому своими руками учебник книга pdf википедия "
                "определение gpt chatgpt гпт чатгпт deepseek дипсик").split()

CAMPAIGNS = [
    {
        "name": "KZ_Воркшоп_12_ноября", "slug": "kz_workshop", "path": "/ru/webinar",
        "minus": [],  # «бесплатно» и «курсы» не минусуем: воркшоп бесплатный и учебный
        "display": "воркшоп-12-ноября",
        "heads": ["Как перестать терять заявки ночью", "ИИ-ассистент для отдела продаж за 45 минут",
                  "12 ноября, 20:00 по Астане", "Бесплатно, онлайн, 12 ноября", "Онлайн-воркшоп 12 ноября"],
        "texts": ["Соберём вживую связку: мессенджер, ИИ-ассистент и CRM. 45 минут, на русском.",
                  "Покажем, как ИИ отвечает клиентам в WhatsApp и сам заносит заявку в CRM.",
                  "Живой разбор на экране и чек-лист в подарок. Запись открыта."],
        "groups": [
            ("Обучение нейросетям", ["обучение нейросетям", "обучение ии и нейросетям"],
             ["Нейросети для бизнеса: бесплатный воркшоп", "Обучение нейросетям на примере продаж"]),
            ("ИИ для бизнеса", ["ии для бизнеса", "нейросети для бизнеса", "нейросеть для бизнеса",
                                "как использовать ии в бизнесе"],
             ["ИИ для бизнеса: бесплатный онлайн-воркшоп", "Как использовать ИИ в бизнесе — вживую"]),
            ("ИИ в продажах", ["ии для отдела продаж", "ии в продажах"],
             ["Бесплатный воркшоп: ИИ в отделе продаж", "ИИ в продажах: покажем на живом примере"]),
        ],
        "sitelinks": [("Записаться бесплатно", "12 ноября, 20:00 по Астане. Ссылку пришлём", "/ru/webinar#register"),
                      ("Вопросы и ответы", "Нужны ли технические знания и что, если не смогу прийти", "/ru/webinar#wb-faq-0"),
                      ("Что покажем", "Связка мессенджер, ИИ-ассистент и CRM за 45 минут", "/ru/webinar#wb-hero")],
        "callouts": ["Бесплатно", "Онлайн", "На русском", "45 минут", "Чек-лист в подарок"],
    },
    {
        "name": "KZ_Услуги_Поиск", "slug": "kz_services", "path": "/ru/departments/sales",
        "minus": ["бесплатно", "бесплатный", "курсы", "курс", "обучение"],
        "display": "ии-для-продаж",
        "heads": ["Внедрение ИИ в отдел продаж", "Пилот за 2–4 недели", "Бесплатный аудит процессов",
                  "Заявки сразу в CRM"],
        "texts": ["ИИ отвечает клиентам в WhatsApp и Telegram и заносит заявки в CRM.",
                  "Ответы клиентам 24/7, заявки в CRM, отчёты. Фиксированная цена пилота.",
                  "Отвечает по вашим ценам и правилам, передаёт сложные вопросы менеджеру."],
        "groups": [
            ("Внедрение ИИ", ["внедрение ии в бизнес", "внедрение ии агентов",
                              "внедрение искусственного интеллекта услуги", "ии агенты для бизнеса"],
             ["Внедрение ИИ в бизнес: пилот за 2–4 недели", "ИИ-агенты для бизнеса под ваши процессы"], "/ru/solutions"),
            ("Чат-боты для бизнеса", ["ии чат бот для бизнеса", "чат бот для бизнеса", "разработка чат бота",
                                      "чат бот для whatsapp"],
             ["Чат-бот для WhatsApp с ИИ", "ИИ-агенты и чат-боты для бизнеса"], None),
            ("Автоматизация продаж", ["автоматизация бизнеса с ии", "ии автоматизация бизнес процессов",
                                      "автоматизация отдела продаж"],
             ["Автоматизация продаж с ИИ", "ИИ заносит заявки в CRM сам"], None),
        ],
        "sitelinks": [("Цены и форматы", "Пилот 2–4 недели по фиксированной цене", "/ru/solutions"),
                      ("Обучение команды", "Практика на задачах вашей команды", "/ru/training"),
                      ("Бесплатный воркшоп", "12 ноября онлайн: ИИ в отделе продаж", "/ru/webinar")],
        "callouts": ["Пилот за 2–4 недели", "Бесплатный аудит", "Фиксированная цена пилота", "Данные на ваших аккаунтах"],
    },
]

SETTINGS_HEAD = ["Кампания", "Тип", "Места показа", "Регион", "Даты", "Расписание (время Астаны, UTC+5)",
                 "Стратегия", "Недельный бюджет, ₽ (предложение)", "Ставки", "Автотаргетинг", "Картинки",
                 "Счётчик Метрики", "Прочее"]
SETTINGS = [
    ["KZ_Воркшоп_12_ноября", "Единая перфоманс-кампания", "Поиск + Рекламная сеть (решение владельца 24.10)",
     REGION, "С запуска до 12 ноября 19:00 по Астане", "Ежедневно 8:00–23:00",
     "Максимум кликов с ручными ставками, недельный бюджет", "≈7 000",
     "Фразы — 30 ₽, автотаргетинг — 0,3 ₽ (как в WinWinBot KZ); в сетях — по прогнозу Коммандера",
     "Только «Целевые» и «Узкие», бренды — «Без упоминания брендов»",
     "Добавить 2 картинки: 1080×607 и 1080×1080 (присланы в чат, на сайт не выкладываются)",
     "113476399 (цели почти не набираются — Метрика только после согласия на cookie)",
     "Название кампании задать вручную (из файла не берётся). «Автоматически применять рекомендации» "
     "и «Персонализация» — выключить, «Мониторинг сайта» — включить, автопометка yclid — включить."],
    ["KZ_Услуги_Поиск", "Единая перфоманс-кампания", "Только поиск (сети, Карты, галереи — выключить)",
     REGION, "С запуска, без даты окончания", "Будни 9:00–20:00",
     "Максимум кликов с ручными ставками, недельный бюджет", "2 000–3 000",
     "Фразы — 30 ₽, автотаргетинг — 0,3 ₽", "Только «Целевые» и «Узкие», бренды — «Без упоминания брендов»",
     "Не нужны (только поиск)", "113476399",
     "Название кампании задать вручную. Рекомендации и персонализация — выключить, мониторинг сайта и "
     "автопометка yclid — включить."],
]


def add_legacy_cols(ws):
    merged = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges]
    for r in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(r))
    ws.insert_cols(9, SHIFT)
    for r1, c1, r2, c2 in merged:
        s = SHIFT if c1 >= 9 else 0
        ws.merge_cells(start_row=r1, start_column=c1 + s, end_row=r2, end_column=c2 + s)
    for i, name in enumerate(LEGACY):
        col = 9 + i
        ws.cell(10, col, name)._style = copy(ws.cell(10, 8)._style)
        ws.merge_cells(start_row=10, start_column=col, end_row=11, end_column=col)


def check(c):
    err = []
    for gname, _, gheads, *_ in c["groups"]:
        heads, texts = ads(c, gheads)
        for h in heads:
            if len(h) > 56 or max(len(w) for w in h.split()) > 22:
                err.append(f"{gname}: заголовок «{h}» ({len(h)})")
        for t in texts:
            if len(t) > 81 or len(PUNCT.findall(t)) > 15 or max(len(w) for w in t.split()) > 23:
                err.append(f"{gname}: текст «{t}» ({len(t)})")
    if len(c["display"]) > 20:
        err.append(f"отображаемая ссылка «{c['display']}»")
    sl = c["sitelinks"]
    if sum(len(t) for t, _, _ in sl[:4]) > 66 or any(len(t) > 30 or len(d) > 60 for t, d, _ in sl):
        err.append("быстрые ссылки длиннее лимита")
    if any(len(x) > 25 for x in c["callouts"]):
        err.append("уточнение длиннее 25")
    phrases = [p for g in c["groups"] for p in g[1]]
    for m in minus(c):
        mw = m.lstrip("-")
        for p in phrases:
            if any(stem(mw) == stem(w) for w in re.findall(r"[a-zа-яё0-9]+", p)):
                err.append(f"минус «{mw}» блокирует фразу «{p}»")
    return err


def stem(w):
    w = w.lower().replace("ё", "е")
    return w if len(w) <= 4 else (w[:len(w) - 2] if len(w) > 6 else w[:4])


def minus(c):
    return ["-" + m for m in sorted(set(MINUS_COMMON + c["minus"]))]


def ads(c, gheads):
    return list(dict.fromkeys(gheads + c["heads"]))[:7], c["texts"][:3]


def fill(wb, c):
    ws = wb["Тексты"]
    if ws.max_row >= FIRST_ROW:
        ws.delete_rows(FIRST_ROW, ws.max_row - FIRST_ROW + 1)
    ws["E7"] = "Единая перфоманс-кампания"
    ws["E8"] = None
    ws["E9"] = " ".join(minus(c))
    ws["A6"] = f"Предложение текстовых блоков для кампании: {c['name']}"
    sl_t = "||".join(t for t, _, _ in c["sitelinks"])
    sl_d = "||".join(d for _, d, _ in c["sitelinks"])
    sl_u = "||".join(utm(c["slug"], u, term=False) for _, _, u in c["sitelinks"])
    row, n = FIRST_ROW, 0
    for num, (gname, phrases, gheads, *rest) in enumerate(c["groups"], 1):
        heads, texts = ads(c, gheads)
        link = utm(c["slug"], (rest[0] if rest and rest[0] else c["path"]))
        for p in phrases:
            n += 1
            ws.cell(row, C["plus"], "-")
            ws.cell(row, C["type"], "Комбинаторное")
            ws.cell(row, C["group"], gname)
            ws.cell(row, C["num"], num)
            ws.cell(row, C["phrase"], p)
            for i, h in enumerate(heads):
                ws.cell(row, C["h"] + i, h)
            for i, t in enumerate(texts):
                ws.cell(row, C["t"] + i, t)
            for i, col in enumerate(LEN_COLS):
                ref = f"{col}{row}"
                ws.cell(row, C["len"] + i, f'=IF({ref}="","",LEN(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('
                        f'SUBSTITUTE(SUBSTITUTE({ref},"!",""),",",""),".",""),";",""),":",""),"""","")))')
            ws.cell(row, C["url"], link)
            ws.cell(row, C["display"], c["display"])
            ws.cell(row, C["region"], REGION)
            ws.cell(row, C["sl_t"], sl_t)
            ws.cell(row, C["sl_d"], sl_d)
            ws.cell(row, C["sl_u"], sl_u)
            ws.cell(row, C["callouts"], "||".join(c["callouts"]))
            row += 1
    return len(c["groups"]), n


def settings(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Настройки"
    ws.append(SETTINGS_HEAD)
    for r in SETTINGS:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Arial")
        cell.fill = PatternFill("solid", fgColor="FFD86B")
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(1, len(SETTINGS_HEAD) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 30
    ws.freeze_panes = "B2"
    wb.save(path)


def main():
    errors = [f"{c['name']}: {e}" for c in CAMPAIGNS for e in check(c)]
    if errors:
        print("\n".join(errors))
        sys.exit(1)
    OUT.mkdir(exist_ok=True)
    wb = load_workbook(TEMPLATE)
    add_legacy_cols(wb["Тексты"])
    for c in CAMPAIGNS:
        g, n = fill(wb, c)
        wb.save(OUT / f"{c['name']}.xlsx")
        print(f"{c['name']}: групп {g}, фраз {n}")
    settings(OUT / "Настройки_кампаний_KZ.xlsx")
    print("Настройки_кампаний_KZ.xlsx")


if __name__ == "__main__":
    main()
